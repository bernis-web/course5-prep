"""Build the course 5 Word report template from its Markdown source."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "docs" / "report-template.md"
DEFAULT_OUTPUT = ROOT / "templates" / "course5-final-report-template.docx"
BODY_FONT = "Microsoft YaHei"
HEADING_FONT = "Microsoft YaHei"
TABLE_HEADER_FILL = "1F4E78"
TABLE_ALT_FILL = "F3F7FA"
TABLE_BORDER = "D9D9D9"


def _set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), TABLE_BORDER)


def _mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))
    _set_run_font(run, BODY_FONT, 9)


def _configure_document(document: Document) -> None:
    core = document.core_properties
    core.title = "跨物种细胞类型整合工具评测"
    core.subject = "课题5最终报告通用模板"
    core.author = "课题5小组"
    core.last_modified_by = "课题5小组"
    core.comments = "由 docs/report-template.md 生成，不包含真实实验结果。"

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    title = document.styles["Title"]
    title.font.name = HEADING_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(24)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = document.styles[style_name]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer)


def _parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def _table_widths(headers: list[str]) -> list[float]:
    if headers == ["方法", "状态", "物种数", "墙钟时间", "峰值内存", "主要问题"]:
        return [0.9, 0.75, 0.7, 0.9, 0.9, 2.15]
    if headers == ["场景", "方法", "GPU", "运行次数", "总时长", "峰值内存", "实际成本"]:
        return [1.0, 0.9, 0.65, 0.75, 0.85, 0.9, 0.9]
    return [6.2 / len(headers)] * len(headers)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    widths = _table_widths(rows[0])
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if column_index == len(values) - 1 and len(values) > 4
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.add_run(value)
            _set_run_font(run, BODY_FONT, 9)
            if row_index == 0:
                _set_cell_shading(cell, TABLE_HEADER_FILL)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                _set_cell_shading(cell, TABLE_ALT_FILL)
    _mark_header_row(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_report_template(source: Path, output: Path) -> Path:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    _configure_document(document)
    first_section = True
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = _parse_table(lines, index)
            _add_table(document, rows)
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph(stripped[2:].strip(), style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("## "):
            if first_section:
                document.add_page_break()
                first_section = False
            document.add_paragraph(stripped[3:].strip(), style="Heading 1")
        elif stripped.startswith("### "):
            document.add_paragraph(stripped[4:].strip(), style="Heading 2")
        else:
            paragraph = document.add_paragraph()
            if first_section:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(stripped.replace("  ", " "))
            _set_run_font(run, BODY_FONT, 10.5)
            if "本模板用于最终 Word 结题报告" in stripped:
                run.font.italic = True
                run.font.color.rgb = RGBColor(89, 89, 89)
        index += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    built = build_report_template(args.source, args.output)
    print(built)


if __name__ == "__main__":
    main()
