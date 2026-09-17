"""Build the Monday course 5 presentation."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "course5-monday-presentation.pptx"

NAVY = RGBColor(0x17, 0x32, 0x4D)
TEAL = RGBColor(0x14, 0x7D, 0x73)
AMBER = RGBColor(0xB7, 0x79, 0x1F)
GRAY = RGBColor(0x51, 0x5A, 0x66)
LIGHT = RGBColor(0xF2, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD6, 0xDE, 0xE4)
DARK_TEXT = RGBColor(0x1D, 0x27, 0x31)
FONT = "Microsoft YaHei"


@dataclass(frozen=True)
class SlideSpec:
    title: str
    subtitle: str | None = None
    bullets: tuple[str, ...] = ()
    table: tuple[tuple[str, ...], ...] | None = None
    process: tuple[str, ...] | None = None
    takeaway: str | None = None
    note: str | None = None


SLIDES: tuple[SlideSpec, ...] = (
    SlideSpec(
        title="跨物种细胞类型整合工具评测",
        subtitle="SATURN vs SAMap vs scVI ｜ 课题5 初步方案汇报",
        bullets=(
            "组员：____、____、____",
            "本次汇报：我们打算怎么做这件事",
            "当前状态：方案已就绪，等课程数据入口",
        ),
        note="开场说明：今天讲的是“怎么做”的方案，还没有实验结果；以后所有结论都会来自我们自己跑出的数据。",
    ),
    SlideSpec(
        title="我们要回答的问题",
        bullets=(
            "背景：三个工具都能把不同物种的细胞“对齐”后比较，但各论文各用各的，没有公平对比",
            "我们的课题：同一批数据、同一套规则，让 SATURN、SAMap、scVI 各跑一遍，按统一标准打分",
            "产出：一张公平的对比表，告诉后来者“什么场景该用哪个工具”",
            "原则：只认我们自己跑出来的数，论文数值和示例数据只能当参考",
        ),
        takeaway="要点：公平比较是课题的核心——同样的数据、同样的规则，只比工具本身。",
        note="这一页回答“进一步理解课题”：核心不是复现论文，而是做一次公平的工具比较。",
    ),
    SlideSpec(
        title="课题的四个关键点",
        table=(
            ("要回答的问题", "关键点", "我们的做法"),
            ("用什么数据", "跨物种单细胞数据，带细胞类型标签", "用课程提供的 HEART 心脏图谱，先拿 2-3 个物种"),
            ("怎么算公平", "三个工具吃同样的输入、同样的规则", "统一细胞、基因、拆分方式和随机种子"),
            ("怎么算好", "既要把物种混匀，也要保住细胞类型差异", "三个工具用同一组指标打分"),
            ("别人能否照做", "换个人也能跑出同样的结果", "每步的输入、参数、结果都存档在仓库里"),
        ),
        note="核心难点是“公平”和“可复现”，这决定了后面所有设计。",
    ),
    SlideSpec(
        title="文献精读 ①：HEART 心脏跨物种图谱",
        bullets=(
            "发表在 Science：构建了动物心脏的跨物种单细胞图谱",
            "规模：27 个物种；其中 16 个物种、约 48 万个细胞做了整合分析",
            "论文做法：用 SATURN 整合物种，用 SAMap 找跨物种对应的细胞类型",
            "关键设置：蛋白质序列嵌入（ESM2B）+ 2000 个宏基因 + 8000 个高变基因",
            "对我们的用处：提供数据来源和标签金标准；但论文结果不能当成我们的结论",
        ),
        takeaway="要点：HEART 提供数据和“标准答案”；论文数字是作者的结果，不是我们的。",
        note="口播稿见讲稿文档第 3 节，60-90 秒版本。",
    ),
    SlideSpec(
        title="文献精读 ②：LBSPA 七鳃鳗脑图谱",
        bullets=(
            "发表在 Science：七鳃鳗全脑图谱，同时有空间数据和单核数据",
            "我们只用单核部分：约 4.2 万个高质量细胞核、209 种细胞簇",
            "论文做过七鳃鳗-小鼠比较：约 6900 个一对一同源基因 + scVI",
            "对我们的用处：单细胞实验做完后，用它做“换个物种跨度再试一次”的补充验证",
            "空间数据暂不进主实验：数据模态不同，混在一起说不清是谁的影响",
        ),
        takeaway="要点：LBSPA 用在最后——单细胞结论出来后，换个物种跨度再验证一次。",
        note="回应老师要求：空间数据是暂缓不是放弃，排在单细胞结论之后。",
    ),
    SlideSpec(
        title="实验总体思路",
        process=("选数据", "小试跑", "统一打分", "扩场景", "写报告"),
        bullets=(
            "选数据：从课程数据里选 2-3 个物种的心脏细胞（心肌细胞、成纤维细胞两类）",
            "小试跑：三个工具各跑通一次，只验证“能不能跑完”，不比名次",
            "统一打分：同一套指标给三个工具打分，重复 3 次取平均",
            "扩场景：加入近缘、远缘、不同类别物种组合，看工具表现怎么变",
            "写报告：结果、耗时、花费、失败原因全部写进报告",
        ),
        note="每一步都有检查点，上一步没通过就不进入下一步。",
    ),
    SlideSpec(
        title="怎么保证比较是公平的",
        bullets=(
            "同一批细胞：三个工具用完全相同的输入，谁也不多给数据",
            "同一张基因对照表：需要基因对应关系的工具，用同一份同源基因表",
            "考试分开出题：轮流留出一个物种当“考题”，用其余物种训练后去猜它的细胞类型（留一物种法）",
            "重复实验：同一实验用 3 个随机种子各跑一遍，报告平均和波动",
            "全程留痕：每次运行的输入、参数、日志、结果都存档，事后可查",
        ),
        takeaway="要点：公平 = 同输入、同规则、考题分开、重复三次、全程留痕。",
        note="“留一物种”是评测设计的核心：避免“用同一批数据既当教练又当考题”。",
    ),
    SlideSpec(
        title="评分指标（每个都配了一句解释）",
        table=(
            ("看什么", "指标", "一句话解释"),
            ("标签猜得准不准", "macro-F1", "被留出物种的细胞类型被猜对的程度"),
            ("物种混得匀不匀", "kBET、iLISI", "整合之后不同物种的细胞是否混在一起"),
            ("类型结构保得好不好", "silhouette", "整合时细胞类型之间的界限有没有被抹掉"),
            ("跑得快不快、贵不贵", "时间、内存、费用", "每次运行的时间和花费都记账"),
        ),
        note="前三类由同一个脚本统一计算，不允许各工具用各自的算法。",
    ),
    SlideSpec(
        title="下一步计划（按检查点推进，不卡死时间）",
        bullets=(
            "第一步（现在）：等老师提供课程数据入口，然后核对数据能读、字段齐全，核对软件环境——这些都不花钱",
            "第二步（全组同意后）：小规模试跑一次，设好保险丝——单卡、限时到点即停、单次花费 10 元封顶",
            "第三步（试跑顺利后）：正式实验、统一评分、写 Word 报告",
            "每一步完成都全组过目；任何一步不通过，就不进入下一步",
        ),
        takeaway="要点：三步走，每步全组过目；上一步不通过，就不进入下一步。",
        note="回应“时间不用太精确”：以检查点推进，不承诺具体小时数。具体限额数字记录在讲稿文档里。",
    ),
    SlideSpec(
        title="外组同学怎么接手（四步就能开始）",
        bullets=(
            "1. 克隆仓库并装依赖：git clone + pip install -r requirements.txt",
            "2. 跑测试确认环境正常：python -m unittest discover -s tests -v",
            "3. 跑一遍示例流程，熟悉输出长什么样：python scripts/smoke_test.py --outdir smoke-output",
            "4. 按仓库里的检查清单（docs/minimal-verification-runbook.md）逐条做核对",
            "每条命令的作用都写在文档里，不需要追问任何“隐藏背景”就能开始",
        ),
        takeaway="要点：四条命令就能开始；命令逐条有解释，无需任何背景知识。",
        note="完整命令清单在讲稿文档第 4 节；现场可以直接演示。",
    ),
    SlideSpec(
        title="分工与互相检查",
        table=(
            ("成员", "主线", "具体负责", "交叉检查"),
            ("____", "找数据", "确认数据路径、物种、标签字段能读、含义清楚", "C"),
            ("____", "管软件和算力", "核对软件版本，控制时长和费用，记录每次运行环境", "A"),
            ("____", "管评分和报告", "统一评分脚本，整理结果表，写报告", "B"),
        ),
        bullets=(
            "主线可以轮换；每个人都要复核别人的一块交付",
            "一条硬规矩：开机花钱的操作，必须全组同意后才执行",
        ),
        note="展示前把成员姓名补齐。",
    ),
    SlideSpec(
        title="可能遇到的问题与对策",
        table=(
            ("可能的问题", "我们的对策"),
            ("课程数据还没拿到", "已在平台全面排查并精确提问老师；回复前不花一分钱"),
            ("软件环境对不上", "进容器核对版本；缺工具只写安装方案，不偷偷装"),
            ("试跑爆内存或超时", "保存现场和日志，缩小数据规模再试，不加钱加卡"),
            ("结果不可比", "输入、种子、版本全部提前写死并存档"),
        ),
        note="每条风险都已经留了证据或写好预案。",
    ),
    SlideSpec(
        title="总结",
        bullets=(
            "方案已完整：数据、方法、评分、分工、费用上限全部明确",
            "已完成：本地流程全部跑通；数据平台已排查两轮，问题已汇总给老师",
            "等待中：课程数据入口——老师回复后即可进入试跑准备",
            "纪律：开机花钱前必须全组同意，跑完立即关机",
            "谢谢大家，欢迎提问",
        ),
        takeaway="要点：方案就绪，等数据入口；花钱之前全组同意。",
        note="结尾回扣老师四项要求：课题理解、文献精读、可执行方案、小组分工均已覆盖。",
    ),
)


def set_text(frame, text: str, size: int, *, bold: bool = False, color: RGBColor = DARK_TEXT) -> None:
    frame.text = text
    run = frame.paragraphs[0].runs[0]
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_box(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int,
    *,
    fill: RGBColor | None = None,
    line: RGBColor | None = LINE,
    bold: bool = False,
    color: RGBColor = DARK_TEXT,
    align=PP_ALIGN.LEFT,
) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.48), Inches(0.13), Inches(0.58))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    add_box(slide, title, 0.82, 0.42, 10.8, 0.55, 27, bold=True, color=NAVY, line=None)
    if subtitle:
        add_box(slide, subtitle, 0.84, 0.97, 10.5, 0.32, 13, color=GRAY, line=None)


def add_bullets(slide, bullets: tuple[str, ...], top: float = 1.48) -> None:
    if not bullets:
        return
    height = min(4.9, 0.42 + len(bullets) * 0.72)
    shape = slide.shapes.add_textbox(Inches(0.82), Inches(top), Inches(11.15), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = item
        run.font.name = FONT
        run.font.size = Pt(17 if len(item) < 44 else 15)
        run.font.color.rgb = DARK_TEXT
        paragraph.space_after = Pt(15)
        paragraph.line_spacing = 1.28


def add_takeaway(slide, text: str) -> None:
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(5.68), Inches(13.333), Inches(1.0))
    band.fill.solid()
    band.fill.fore_color.rgb = LIGHT
    band.line.fill.background()
    shape = slide.shapes.add_textbox(Inches(0.82), Inches(5.82), Inches(11.7), Inches(0.72))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = NAVY


def add_table(slide, rows: tuple[tuple[str, ...], ...]) -> None:
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.75), Inches(1.48), Inches(10.9), Inches(3.9))
    table = shape.table
    widths = _table_widths(len(rows[0]))
    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Inches(width)
    for row_index, values in enumerate(rows):
        table.rows[row_index].height = Inches(0.62 if row_index else 0.48)
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if row_index == 0 else (LIGHT if row_index % 2 == 0 else WHITE)
            frame = cell.text_frame
            frame.clear()
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(12 if row_index == 0 else 11)
            run.font.bold = row_index == 0
            run.font.color.rgb = WHITE if row_index == 0 else DARK_TEXT


def _table_widths(columns: int) -> tuple[float, ...]:
    if columns == 3:
        return (1.75, 4.35, 4.8)
    if columns == 4:
        return (1.4, 2.7, 5.1, 1.7)
    return (10.9 / columns,) * columns


def add_process(slide, steps: tuple[str, ...]) -> None:
    left = 0.72
    top = 1.52
    width = 1.42
    height = 0.58
    for index, step in enumerate(steps):
        x = left + index * (width + 0.16)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top), Inches(width), Inches(height))
        shape.adjustments[0] = 0.16
        shape.fill.solid()
        shape.fill.fore_color.rgb = TEAL if index in (1, 4, 6) else NAVY
        shape.line.fill.background()
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.03)
        frame.margin_right = Inches(0.03)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = step
        run.font.name = FONT
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = WHITE
        if index < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(x + width + 0.025),
                Inches(top + 0.20),
                Inches(0.11),
                Inches(0.18),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = AMBER
            arrow.line.fill.background()


def add_footer(slide, number: int) -> None:
    add_box(slide, "课题5 | 跨物种单细胞整合工具评测", 0.75, 7.02, 5.0, 0.25, 9, color=GRAY, line=None)
    add_box(slide, str(number), 11.7, 7.02, 0.45, 0.25, 9, color=GRAY, line=None, align=PP_ALIGN.RIGHT)


def build() -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    for index, spec in enumerate(SLIDES, start=1):
        slide = presentation.slides.add_slide(blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = WHITE
        if index == 1:
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.22))
            accent.fill.solid()
            accent.fill.fore_color.rgb = TEAL
            accent.line.fill.background()
            add_box(slide, spec.title, 0.85, 1.95, 10.6, 1.0, 34, bold=True, color=NAVY, line=None)
            if spec.subtitle:
                add_box(slide, spec.subtitle, 0.88, 2.95, 9.8, 0.45, 18, color=TEAL, line=None)
            add_bullets(slide, spec.bullets, top=3.75)
        else:
            add_title(slide, spec.title, spec.subtitle)
            if spec.process:
                add_process(slide, spec.process)
            if spec.table:
                add_table(slide, spec.table)
            if spec.takeaway:
                add_takeaway(slide, spec.takeaway)
            bullets_top = 2.35 if spec.process else (1.75 if spec.takeaway else 1.48)
            if not spec.table:
                add_bullets(slide, spec.bullets, top=bullets_top)
            elif spec.bullets:
                add_bullets(slide, spec.bullets, top=4.95)
        if spec.note:
            slide.notes_slide.notes_text_frame.text = spec.note
        add_footer(slide, index)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    return OUTPUT


def main() -> None:
    print(build())


if __name__ == "__main__":
    main()
